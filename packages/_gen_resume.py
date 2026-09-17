# -*- coding: utf-8 -*-
"""生成极限精简版简历 .docx（保留全部技能点，单行短句化）"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

def set_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

def heading(text, size=13, before=6, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_font(r, size=size, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '8'), ('w:space', '1'), ('w:color', '1F4E79')):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p

def body(text, bold=False, size=10.5, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p

def bullet(text, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run("· " + text)
    set_font(r)
    return p

# ============ 标题 ============
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_after = Pt(2)
r = t.add_run("杨顺发"); set_font(r, size=20, bold=True)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(8)
r = sub.add_run("求职意向：AI 应用开发工程师（长沙）"); set_font(r, size=12, bold=True, color=RGBColor(0xC0,0x50,0x00))

# ============ 基本信息 ============
heading("基本信息")
body("1997.02　|　本科·湖南文理学院芙蓉学院·计算机科学与技术　|　四年以上经验　|　全职")

# ============ 专业技能 ============
heading("专业技能")
body("【AI 应用开发】", bold=True)
bullet("Agent：LangGraph 实战 ReAct；理解 CoT / ToT / Plan-and-Execute，按任务复杂度选型编排；")
bullet("Prompt / Tokenizer / Memory：防幻觉与 SOP 提示词；token 预算控制上下文与分块；分层 Memory（临时上下文 / checkpoint 会话 / 向量库长期）；")
bullet("RAG：ChromaDB 入库；Hybrid Search（向量 + BM25/jieba）→ RRF 融合 → gte-rerank-v2 Rerank 精排；引用溯源防幻觉；")
bullet("流式与后端：astream_events + LCEL 的 SSE（思考/工具/token 事件分离）；FastAPI 编排、Spring Boot 网关（调度/限流熔断/计量）、NestJS（JWT/权限/审计）。")
body("【语言与工程】", bold=True)
bullet("Python（FastAPI）、TypeScript；LangChain/LangGraph、向量库、Embedding/ReRank；MySQL、SQLite、ChromaDB、Linux 部署。")
body("【前端（辅助）】", bold=True)
bullet("Vue2/Vue3 + TS：Vue Router、Vuex/Pinia、axios 二次封装、Webpack/Vite、ES6 异步（Promise/async-await/Event Loop）、Element Plus；SSE 流式渲染对话界面。")

# ============ 工作经验 ============
heading("工作经验")
body("创智和宇信息技术有限公司（长沙）　2020.10–至今　AI 应用开发工程师", bold=True)
body("主导 Agent 编排、RAG、模型网关的大模型应用工程化，并负责和宇云 PaaS 前端架构；参与智慧医保、分布式数据库云平台。")
body("金联达智能包装科技有限公司（深圳）　2018.10–2020.09　前端开发工程师", bold=True)
body("包装宝平台 Web/H5/小程序开发。")

# ============ 项目经验 ============
heading("项目经验")

body("对话式 AI 智能助手平台（Agent + RAG + 模型网关）　2025–至今　8 人", bold=True)
bullet("链路：浏览器 → NestJS（鉴权）→ ai-service（LangGraph + SSE）→ 模型网关（唯一模型出口）→ 大模型；")
bullet("Agent：ReAct 循环，checkpoint + messages 双轨持久化，系统提示临时前置防上下文膨胀；")
bullet("SSE：astream_events 拆 reasoning/tool/token 事件，子图命名空间过滤，思考正文分离、工具步骤可视化；")
bullet("RAG：Agent 经 search_docs 工具自主调外部知识库，编排与知识库解耦；")
bullet("网关与权限：多上游调度、限流熔断、调用计量看板；NestJS 两维授权（职级批量 + 个人覆盖）+ 审计。")

body("企业级私有文档问答系统（RAG，独立完成）　2025–至今", bold=True)
bullet("入库：PDF/TXT 解析（编码回退）→ 递归切分（500/50）→ 向量化入库 + BM25 索引；")
bullet("检索：向量 + BM25 双路召回 → RRF 融合（k=60）→ gte-rerank-v2 精排，失败降级 RRF；")
bullet("生成与流式：防幻觉提示 + [资料N] 引用溯源，无召回不调 LLM；SSE 先推 sources 再逐 token 输出，asyncio.to_thread 防阻塞；")
bullet("工程：JWT 鉴权，第三方经 X-Service-Key 免登录调用检索接口。")

body("和宇云 PaaS 平台、分布式数据库　2020.10–2025　30 人", bold=True)
bullet("分布式权限模块 + 动态路由，按公司/个人管控访问；主机监控告警 ECharts 可视化；前端选型与组件化规范。")

body("贵联控股 - 包装宝平台系统　2018.10–2020.09　8 人", bold=True)
bullet("Vue SPA 多端交付（Web/H5/小程序）：路由/状态管理/组件化与接口二次封装。")

# ============ 在校经历 ============
heading("在校经历")
body("班长、学生会/院团委办公室主任；校级优秀团干部、三好学生，两次一等奖学金。")

# ============ 自我评价 ============
heading("自我评价")
body("独立完成 Agent 编排、企业级 RAG 与模型网关的完整落地，工程上注重分层解耦、可靠性兜底与可观测性；"
     "学习与问题驱动，习惯以文档沉淀方案，能将大模型能力封装为可交付产品。")

out = r"d:\selfFile\python\vue3\v3agent\packages\AI应用开发工程师(AI模块版) trea版本 .docx"
doc.save(out)
print("SAVED:", out)